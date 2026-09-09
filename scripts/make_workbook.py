"""Companion workbook: frozen scenario outputs and explicit formula examples."""
import json
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.comments import Comment

def make_workbook(root):
    wb=Workbook();ws=wb.active;ws.title='Read me'
    rows=[['SSM - scenario data','2026-09-08'],['Status','Synthetic mechanism study; not a forecast or estimated EU model.'],
          ['Reproduction','python scripts/reproduce_all.py'],['Units','Fractions are dimensionless unless column name specifies otherwise.'],
          ['EU scaling','2025 GDP, April 2026 EDP vintage; rounded January 2026 resident population.'],
          ['Coverage','Full trajectories, horizon outcomes, robustness, sensitivity and analytical envelope in repository CSVs.'],
          ['Important','Near-target inflation in adaptive cases relies on residual taxes. A completed run is not necessarily price-stable.'],
          ['Empirical basis','Official anchors only. Household distributions and behavioural parameters are synthetic assumptions.'],['Revision','v0.2.0: 549 macro runs including repeated controls; matched payouts are identical, legal wealth differs.']]
    for row in rows:ws.append(row)
    for filename,title in [('horizons.csv','Horizons'),('robustness.csv','Robustness'),('sensitivity.csv','Sensitivity'),
                           ('policy_comparison_summary.csv','Policy comparison'),('policy_comparison_robustness.csv','Policy robustness'),('matched_equivalence_checks.csv','Equivalence checks'),('ownership_incidence.csv','Incidence'),('micro_examples.csv','Micro examples'),('dilution_arithmetic.csv','Equity arithmetic'),('frontier.csv','Frontier')]:
        df=pd.read_csv(root/'results/tables'/filename)
        preferred={
            'Horizons':['scenario','code','horizon','completed','real_gdp','usownership_share','mvi_gdp','tax_gdp','gini_wealth','max_inflation'],
            'Policy comparison':['scenario','code','year','real_gdp','usownership_share','mvi_gdp','tax_gdp','participant_income_gdp','gini_wealth','max_inflation','pass_through_levy_gdp','total_recorded_tax_gdp'],
            'Policy robustness':['experiment','code','year','status','real_gdp','mvi_gdp','tax_gdp','usownership_share','gini_wealth','max_inflation'],
            'Robustness':['experiment','year','status','real_gdp','mvi_gdp','tax_gdp','usownership_share','gini_wealth','max_inflation','sunset']
        }.get(title,[])
        if preferred:df=df[preferred+[c for c in df if c not in preferred]]
        s=wb.create_sheet(title);s.append(list(df.columns))
        for row in df.itertuples(index=False,name=None):
            s.append([None if pd.isna(v) else v for v in row])
    t=pd.read_csv(root/'results/tables/trajectories.csv')
    fields=['scenario','code','year','real_gdp','mvi_gdp','tax_gdp','usownership_share','gini_wealth','inflation','fiscal_issuance_gdp',
            'net_seigniorage_gdp','burn_gdp','labour_share','use_dividend_gdp','mvi_real_eur_resident_year']
    s=wb.create_sheet('Plotted trajectories');s.append(fields)
    for row in t[fields].itertuples(index=False,name=None):s.append(row)
    s=wb.create_sheet('Parameters');s.append(['Parameter','Value','Provenance'])
    import yaml
    for key,value in yaml.safe_load((root/'configs/baseline.yaml').read_text()).items():
        s.append([key,value,'Scenario assumption / numerical design, not empirical estimate'])
    s=wb.create_sheet('Official anchors');s.append(['Variable','Value','Source'])
    a=json.loads((root/'data/processed/anchors.json').read_text())
    for key in ['eu_gdp_eur_2025','eu_government_expenditure_gdp_2025','eu_government_revenue_gdp_2025','eu_government_debt_gdp_2025',
                'ea20_gdp_eur_2025','population_scaling_2026']:
        url=a['population_source'] if 'population' in key else a['gdp_source']
        s.append([key,a[key],url]);s.cell(s.max_row,2).comment=Comment(url,'Source')
    s=wb.create_sheet('Exact identities');s.append(['Input or derived quantity','Value','Interpretation'])
    for row in [['Real growth',.04,'Assumption'],['Target inflation',.02,'Assumption'],['Gross issuance / opening money',.12,'Assumption'],
                ['Exact nominal growth','=(1+B2)*(1+B3)-1','Derived, not the sum approximation'],
                ['Required burn / opening money','=B4-B5','Controller solution, not an estimated policy effect'],
                ['Monthly decay',.01,'Assumption'],['Annual decay','=1-(1-B7)^12','A cohort decay, not burn/GDP'],
                ['Annual equity issue / old equity',.02,'Assumption'],['Years',50,'Assumption'],
                ['Cumulative user share','=1-(1+B9)^(-B10)','No sales; all new shares allocated to users']]:s.append(row)
    for sh in wb:
        sh.freeze_panes='A2';sh.sheet_view.showGridLines=False;sh.auto_filter.ref=sh.dimensions
        for cell in sh[1]:
            cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='203864');cell.alignment=Alignment(wrap_text=True)
        sh.row_dimensions[1].height=32
        for col in sh.columns:
            letter=col[0].column_letter
            size=max(min(max(len(str(c.value or '')) for c in list(col)[:100])+2,50),14)
            sh.column_dimensions[letter].width=size
        for row in sh.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,float):c.number_format='0.0000;[Red](0.0000);-'
                if c.data_type=='f':c.font=Font(color='000000')
                elif isinstance(c.value,(int,float)):c.font=Font(color='008000')
    wb['Read me'].column_dimensions['B'].width=100
    for row in wb['Read me']:row[1].alignment=Alignment(wrap_text=True,vertical='top');wb['Read me'].row_dimensions[row[0].row].height=34
    for row in wb['Exact identities'].iter_rows(min_row=2):
        value=row[1]
        value.number_format='0.00%;[Red](0.00%);-'
        if value.data_type!='f':
            value.font=Font(color='0000FF')
            value.comment=Comment('Illustrative assumption; not an observed estimate.','Provenance')
    wb['Exact identities']['B10'].number_format='0'
    for row in wb['Parameters'].iter_rows(min_row=2):
        row[1].font=Font(color='0000FF')
        row[1].comment=Comment('Scenario assumption recorded in configs/baseline.yaml.','Provenance')
    for sh in wb:
        for cell in sh[1]:
            if (str(cell.value).endswith('_gdp') and str(cell.value) not in ['real_gdp','nominal_gdp']) or str(cell.value) in ['inflation','max_inflation','usownership_share','labour_share','demurrage','annual_dilution','user_share','dilution_cap','income_floor','annual_issuance','legacy_pv_loss','required_total_value_growth_for_legacy_wealth']:
                for colcell in sh.iter_rows(min_row=2,min_col=cell.column,max_col=cell.column):
                    if isinstance(colcell[0].value,(float,int)):colcell[0].number_format='0.00%;[Red](0.00%);-'
    wb['Official anchors']['B2'].number_format='#,##0'
    wb['Official anchors']['B6'].number_format='#,##0'
    wb['Official anchors']['B7'].number_format='#,##0'
    for cell in ['B3','B4','B5']:wb['Official anchors'][cell].number_format='0.0%'
    wb.save(root/'results/scenario_companion.xlsx')
    cache_formula_values(root/'results/scenario_companion.xlsx')


def cache_formula_values(path):
    """Populate cached values for the four documented formula cells, retaining formulas.

    openpyxl does not calculate formulas. These identities are evaluated here from
    the same input cells and inserted into the OOXML cache; Excel can recalculate
    them normally. Unsupported formula changes fail rather than being guessed.
    """
    from openpyxl import load_workbook
    from zipfile import ZipFile, ZIP_DEFLATED
    import xml.etree.ElementTree as ET
    from io import BytesIO
    book = load_workbook(path)
    sheet = book['Exact identities']
    expected = {'B5':'=(1+B2)*(1+B3)-1','B6':'=B4-B5','B8':'=1-(1-B7)^12','B11':'=1-(1+B9)^(-B10)'}
    for key, formula in expected.items():
        if sheet[key].value != formula:
            raise ValueError('Unexpected formula in '+key)
    growth = (1+sheet['B2'].value)*(1+sheet['B3'].value)-1
    values = {'B5':growth,'B6':sheet['B4'].value-growth,
              'B8':1-(1-sheet['B7'].value)**12,'B11':1-(1+sheet['B9'].value)**(-sheet['B10'].value)}
    sheet_index = book.sheetnames.index('Exact identities')+1
    target = f'xl/worksheets/sheet{sheet_index}.xml'
    ns = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    output = BytesIO()
    with ZipFile(path) as source, ZipFile(output,'w',ZIP_DEFLATED) as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == target:
                root = ET.fromstring(data)
                for cell in root.iter(ns+'c'):
                    ref = cell.attrib.get('r')
                    if ref in values:
                        val = cell.find(ns+'v')
                        if val is None: val = ET.SubElement(cell,ns+'v')
                        val.text = repr(values[ref])
                data = ET.tostring(root,encoding='utf-8',xml_declaration=True)
            dest.writestr(item,data)
    path.write_bytes(output.getvalue())
    cached = load_workbook(path,data_only=True)['Exact identities']
    for key, value in values.items():
        if abs(cached[key].value-value)>1e-12:
            raise AssertionError('Incorrect formula cache: '+key)
