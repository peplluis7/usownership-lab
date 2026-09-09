"""Checks on the included full research outputs; rerun default design first."""
from pathlib import Path
import json
import pandas as pd
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]

def test_main_output_integrity():
    frame=pd.read_csv(ROOT/'results/tables/trajectories.csv')
    assert len(frame.groupby(['scenario','code']))==30
    assert frame.accounting_residual.abs().max()<1e-10
    assert frame.goods_residual.abs().max()<1e-10
    assert frame.equity_residual.abs().max()<1e-10

def test_frozen_research_counts():
    assert len(pd.read_csv(ROOT/'results/tables/robustness.csv'))==36
    sensitivity=pd.read_csv(ROOT/'results/tables/sensitivity.csv')
    assert len(sensitivity)==384
    assert sensitivity.groupby('code').sunset.sum().to_dict()=={'B2':0,'B5':4,'B6':6}

def test_manuscript_numeric_macros():
    h=pd.read_csv(ROOT/'results/tables/horizons.csv')
    row=h[(h.scenario=='medium')&(h.code=='B6')&(h.horizon==50)].iloc[0]
    macros=json.loads((ROOT/'results/key_findings.json').read_text())
    assert macros['ResultPSixMviGdp']==f'{100*row.mvi_gdp:.2f}'
    assert macros['ResultPSixGiniWealth']==f'{row.gini_wealth:.3f}'
    assert macros['SensitivityCount']=='128'

def test_workbook_formulas_and_caches():
    path=ROOT/'results/scenario_companion.xlsx'
    f=load_workbook(path,data_only=False)['Exact identities']
    v=load_workbook(path,data_only=True)['Exact identities']
    for key,expected in {'B5':.0608,'B6':.0592,'B8':1-.99**12,'B11':1-1.02**-50}.items():
        assert f[key].data_type=='f'
        assert abs(v[key].value-expected)<1e-12

def test_comparative_output_counts_and_matching():
    frame=pd.read_csv(ROOT/'results/tables/policy_comparison_paths.csv')
    assert len(frame.groupby(['scenario','code']))==27
    assert len(pd.read_csv(ROOT/'results/tables/policy_comparison_robustness.csv'))==72
    checks=pd.read_csv(ROOT/'results/tables/matched_equivalence_checks.csv')
    assert len(checks)==48
    assert checks.max_absolute_error.max()<1e-10

def test_comparison_workbook_tabs():
    book=load_workbook(ROOT/'results/scenario_companion.xlsx',read_only=True)
    assert {'Policy comparison','Policy robustness','Equivalence checks','Incidence'}.issubset(book.sheetnames)
    book.close()

def test_reference_keys_and_package_version():
    import ssm
    assert ssm.__version__=='0.3.0'
    rows=json.loads((ROOT/'docs/source_ledger.json').read_text())
    keys=[r['key'] for r in rows]
    assert len(keys)==51 and len(set(keys))==51
